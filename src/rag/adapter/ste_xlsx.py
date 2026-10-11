from collections.abc import Iterator
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

from rag.domain.suggestions import SteItem

ITEMS_SHEET = 0
CHARACTERISTICS_SHEET = 1
SEPARATOR = " | "


def _rows(path: Path, sheet: int) -> Iterator[dict[str, Any]]:
    book = load_workbook(path, read_only=True)
    try:
        rows = book.worksheets[sheet].iter_rows(values_only=True)
        header = [str(h).strip() if h else "" for h in next(rows)]
        for row in rows:
            yield dict(zip(header, row))
    finally:
        book.close()


def _split(part: str, names: set[str]) -> tuple[str, str] | None:
    words = part.split()
    for k in range(len(words) - 1, 0, -1):
        if " ".join(words[:k]).lower() in names:
            return " ".join(words[:k]), " ".join(words[k:])
    return None


class XlsxSteSource:
    def __init__(self, directory: Path) -> None:
        self._directory = directory

    def _files(self) -> list[Path]:
        files = sorted(self._directory.glob("*.xlsx"))
        if not files:
            raise FileNotFoundError(f"в {self._directory} нет xlsx")
        return files

    def items(self) -> Iterator[SteItem]:
        files = self._files()
        names: set[str] = set()
        categorical: set[str] = set()
        for f in files:
            for row in _rows(f, CHARACTERISTICS_SHEET):
                name = str(row.get("Название характеристики") or "").strip().lower()
                names.add(name)
                kind = str(row.get("Тип (базовая/эталонная/категореобразующая)") or "")
                if kind.startswith("Категор"):
                    categorical.add(name)
        for f in files:
            for row in _rows(f, ITEMS_SHEET):
                title, text = row.get("Наименование СТЕ"), row.get("Характеристики")
                if not isinstance(title, str) or not isinstance(text, str):
                    continue
                parts = (_split(p.strip(), names) for p in text.split(SEPARATOR))
                chars = {
                    n: v for n, v in filter(None, parts) if n.lower() not in categorical
                }
                yield SteItem(title, chars)
