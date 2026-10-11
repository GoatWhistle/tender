import asyncio
import pickle
import re
import sys
import tempfile
import time
from collections import Counter, defaultdict
from dataclasses import dataclass
from functools import cache
from pathlib import Path

from claude_agent_sdk import ClaudeAgentOptions, ResultMessage, query
from openpyxl import load_workbook
from sklearn.feature_extraction.text import TfidfVectorizer

from rag.index import DATA
from rag.models import Position
from rag.spec import card

STE = DATA / "ste.pkl"
STE_SOURCE = DATA / "ste"
ANALOGS = 300
MIN_SCORE = 0.2
MIN_SHARE = 0.04
TOP = 25
LIMIT = 6
SKIP = re.compile(
    r"страна|производител|бренд|торгов|модель|артикул|штрих|видео|упаковк|гарант|"
    r"серия|коллекц|комплектац|ндс|цена|ссылк|сертиф|^вид товаров|^вид продукции",
    re.IGNORECASE,
)

SYSTEM_PROMPT = (
    "Ты помогаешь заполнить карточку товара для госзакупки по 44-ФЗ. Предложи до 6 "
    "дополнительных характеристик, которых нет в КТРУ и которые нужны, чтобы описать "
    "конкретный товар и проверить его при приёмке (п. 5 Правил использования КТРУ, ПП РФ № 145).\n"
    "Источники:\n"
    "- аналоги — характеристики похожих товаров каталога СТЕ с долей товаров, где они указаны. "
    "Сначала реши, тот же ли это вид товара по назначению (analogs_relevant). Если аналоги — "
    "другой товар (например, строительный инструмент вместо медицинского), не бери их;\n"
    "- собственные знания о товаре, если аналоги не подходят или их мало.\n"
    "Правила для характеристики:\n"
    "- не дублирует и не переименовывает характеристику КТРУ и не повторяет название позиции;\n"
    "- объективна и проверяется при приёмке: измерение, материал, наличие, конструкция;\n"
    "- без товарного знака, модели, производителя, страны, артикула, упаковки (ст. 33 44-ФЗ);\n"
    "- важна заказчику: влияет на назначение, безопасность, совместимость или срок службы. "
    "Цвет и декор — только если они функциональны;\n"
    "- название в стиле КТРУ: существительное с уточнением, без значения и единицы в названии; "
    "синонимы аналогов склей в одно название.\n"
    "value — значение для этого товара только из описания поставщика, без единицы измерения; "
    "evidence — дословная цитата из описания. Если в описании значения нет, value и evidence — "
    "пустые строки, не угадывай. options — до 6 типичных значений без единицы (из аналогов или "
    "общие), unit — единица измерения или пустая строка. source: «СТЕ», если взята из аналогов, "
    "иначе «ИИ». why — одно короткое предложение, зачем она заказчику. Сортируй по важности."
)

SCHEMA = {
    "type": "object",
    "properties": {
        "analogs_relevant": {"type": "boolean"},
        "analogs_comment": {"type": "string"},
        "characteristics": {
            "type": "array",
            "maxItems": LIMIT,
            "items": {
                "type": "object",
                "properties": {
                    "name": {"type": "string"},
                    "unit": {"type": "string"},
                    "options": {"type": "array", "items": {"type": "string"}},
                    "value": {"type": "string"},
                    "evidence": {"type": "string"},
                    "source": {"type": "string", "enum": ["СТЕ", "ИИ"]},
                    "why": {"type": "string"},
                },
                "required": ["name", "unit", "options", "value", "evidence", "source", "why"],
                "additionalProperties": False,
            },
        },
    },
    "required": ["analogs_relevant", "analogs_comment", "characteristics"],
    "additionalProperties": False,
}


@dataclass(frozen=True, slots=True)
class Stat:
    name: str
    share: float
    values: tuple[str, ...]


@dataclass(frozen=True, slots=True)
class Suggestion:
    name: str
    unit: str
    options: tuple[str, ...]
    value: str
    evidence: str
    source: str
    share: float | None
    why: str


def norm(name: str) -> str:
    return re.sub(r"[^а-яёa-z0-9]+", " ", name.lower()).strip()


def stem(name: str) -> frozenset[str]:
    return frozenset(w[:5] for w in norm(name).split() if len(w) > 2)


def _split(part: str, names: set[str]) -> tuple[str, str] | None:
    words = part.split()
    for k in range(len(words) - 1, 0, -1):
        if " ".join(words[:k]).lower() in names:
            return " ".join(words[:k]), " ".join(words[k:])
    return None


def _rows(path: Path, sheet: int):
    book = load_workbook(path, read_only=True)
    rows = book.worksheets[sheet].iter_rows(values_only=True)
    header = [str(h).strip() if h else "" for h in next(rows)]
    for row in rows:
        yield dict(zip(header, row))
    book.close()


def build(source: Path = STE_SOURCE) -> None:
    start = time.perf_counter()
    files = sorted(source.glob("*.xlsx"))
    if not files:
        raise SystemExit(f"в {source} нет xlsx")
    names: set[str] = set()
    categorical: set[str] = set()
    for f in files:
        for row in _rows(f, 1):
            name = str(row.get("Название характеристики") or "").strip().lower()
            names.add(name)
            if str(row.get("Тип (базовая/эталонная/категореобразующая)") or "").startswith("Категор"):
                categorical.add(name)
    items = []
    for f in files:
        for row in _rows(f, 0):
            title, text = row.get("Наименование СТЕ"), row.get("Характеристики")
            if not isinstance(title, str) or not isinstance(text, str):
                continue
            parts = (_split(p.strip(), names) for p in text.split(" | "))
            chars = {n: v for n, v in filter(None, parts) if n.lower() not in categorical}
            items.append((title, chars))
        print(f"{f.name}: {len(items)} СТЕ, {time.perf_counter() - start:.0f}s", flush=True)
    vectorizer = TfidfVectorizer(analyzer="char_wb", ngram_range=(3, 5), min_df=2, max_features=400_000)
    matrix = vectorizer.fit_transform(t.lower() for t, _ in items)
    STE.write_bytes(pickle.dumps({"items": items, "vectorizer": vectorizer, "matrix": matrix}))
    print(f"{STE}: {len(items)} СТЕ, {time.perf_counter() - start:.0f}s")


def build_index() -> None:
    build(Path(sys.argv[1]) if len(sys.argv) > 1 else STE_SOURCE)


@cache
def load() -> dict | None:
    return pickle.loads(STE.read_bytes()) if STE.exists() else None


def analogs(text: str) -> tuple[list[str], list[Stat]]:
    ste = load()
    if ste is None:
        return [], []
    scores = (ste["matrix"] @ ste["vectorizer"].transform([text.lower()]).T).toarray().ravel()
    top = [i for i in scores.argsort()[::-1][:ANALOGS] if scores[i] > MIN_SCORE]
    if not top:
        return [], []
    count: Counter[str] = Counter()
    values: defaultdict[str, Counter[str]] = defaultdict(Counter)
    label: dict[str, str] = {}
    for i in top:
        for name, value in ste["items"][i][1].items():
            if SKIP.search(name):
                continue
            key = norm(name)
            label.setdefault(key, name)
            count[key] += 1
            values[key][value.lower()] += 1
    stats = [
        Stat(label[k], n / len(top), tuple(v for v, _ in values[k].most_common(5)))
        for k, n in count.most_common(TOP)
        if n / len(top) >= MIN_SHARE
    ]
    return [ste["items"][i][0] for i in top], stats


def build_prompt(p: Position, description: str, titles: list[str], stats: list[Stat]) -> str:
    lines = [card(p, {}, []), "", f"Описание поставщика:\n{description}", ""]
    if titles:
        lines.append(f"Аналоги СТЕ ({len(titles)} шт.), примеры названий:")
        lines.extend(f"- {t}" for t in titles[:12])
        lines.append("Характеристики аналогов (доля товаров: примеры значений):")
        lines.extend(f"- {s.name} ({s.share:.0%}): {'; '.join(s.values)}" for s in stats)
    else:
        lines.append("Аналогов СТЕ не найдено.")
    return "\n".join(lines)


def _strip_unit(value: str, unit: str) -> str:
    value = value.strip()
    if unit and value.lower().endswith(unit.lower()):
        value = value[: -len(unit)].strip()
    return value


def clean(p: Position, description: str, out: dict, stats: list[Stat]) -> list[Suggestion]:
    taken = {norm(c.name) for c in p.characteristics} | {norm(p.name)}
    known = [(stem(s.name), s.share) for s in stats] if out["analogs_relevant"] else []
    found = []
    for c in out["characteristics"]:
        name, unit = c["name"].strip(), c["unit"].strip()
        if not name or norm(name) in taken:
            continue
        taken.add(norm(name))
        own = stem(name)
        share = max(
            (s for k, s in known if own and len(own & k) / len(own | k) >= 0.5), default=None
        )
        value, evidence = c["value"], c["evidence"].strip()
        if not evidence or evidence.lower() not in description.lower():
            value = evidence = ""
        options = tuple(dict.fromkeys(o for o in (_strip_unit(o, unit) for o in c["options"]) if o))
        found.append(
            Suggestion(
                name=name,
                unit=unit,
                options=options,
                value=_strip_unit(value, unit),
                evidence=evidence,
                source="СТЕ" if share is not None and c["source"] == "СТЕ" else "ИИ",
                share=share if c["source"] == "СТЕ" else None,
                why=c["why"].strip(),
            )
        )
    return found


async def suggest(p: Position, description: str, model: str = "sonnet") -> list[Suggestion]:
    titles, stats = await asyncio.to_thread(analogs, f"{p.name} {description}")
    result: ResultMessage | None = None
    with tempfile.TemporaryDirectory(prefix="rag-agent-") as cwd:
        options = ClaudeAgentOptions(
            model=model,
            effort="low",
            system_prompt=SYSTEM_PROMPT,
            setting_sources=[],
            tools=[],
            cwd=cwd,
            output_format={"type": "json_schema", "schema": SCHEMA},
        )
        async for message in query(prompt=build_prompt(p, description, titles, stats), options=options):
            if isinstance(message, ResultMessage):
                result = message
    if result is None:
        raise RuntimeError("нет ResultMessage")
    if result.is_error or result.structured_output is None:
        raise RuntimeError(result.errors or result.result)
    return clean(p, description, result.structured_output, stats)
